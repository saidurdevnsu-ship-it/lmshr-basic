import shutil
import tempfile
from datetime import date
from io import BytesIO, StringIO
from unittest import mock

import pymupdf
from docx import Document
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from audit.models import AuditEvent
from notifications.models import Notification

from . import ai, workflow
from .forms import GenerateQuestionsForm
from .models import (
    AcademicEvent, BlueprintItem, FinalPaper, Grade, KnowledgeSource, Question, QuestionBlueprint, Subject,
)
from .reminders import ensure_reminders

MEDIA = tempfile.mkdtemp()


def claude_reply(n):
    return {"questions": [
        {"section": "whatever", "question_type": "short", "text": f"Question {i}?", "options": [], "answer": "A",
         "explanation": "", "marks": 99, "source_reference": "Ch 1"}
        for i in range(n)
    ]}


@override_settings(MEDIA_ROOT=MEDIA)
class BlueprintTestCase(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        User = get_user_model()
        self.hod = User.objects.create_user("hod", password="Pw-12345!", is_staff=True)
        self.hod.user_permissions.add(Permission.objects.get(codename="review_blueprint"))
        self.hod = User.objects.get(pk=self.hod.pk)  # reload to clear the permission cache
        self.controller = User.objects.create_superuser("controller", "c@example.com", "Pw-12345!")
        subject, grade = Subject.objects.create(name="Science"), Grade.objects.create(name="Class 8")
        KnowledgeSource.objects.create(subject=subject, grade=grade, title="Book",
                                       file=SimpleUploadedFile("book.txt", b"Plants make food."))
        self.bp = QuestionBlueprint.objects.create(
            title="Science mid-term", subject=subject, grade=grade, assigned_to=self.hod, due_date=date(2026, 12, 1),
            status=QuestionBlueprint.Status.APPROVED,
        )
        BlueprintItem.objects.create(blueprint=self.bp, order=1, section="A", question_type="mcq", count=2, marks_each=1)
        BlueprintItem.objects.create(blueprint=self.bp, order=2, section="B", question_type="short", count=1, marks_each=5)


class FinalPaperTests(BlueprintTestCase):
    def test_marks_sections_and_types_follow_the_blueprint(self):
        with mock.patch.object(ai, "_ask_claude", return_value=claude_reply(3)):
            paper = ai.generate_final_paper(self.bp, self.controller)
        questions = list(paper.questions.all())
        self.assertEqual([q.section for q in questions], ["A", "A", "B"])
        self.assertEqual([q.question_type for q in questions], ["mcq", "mcq", "short"])
        self.assertEqual(paper.total_marks, 7)

    def test_wrong_number_of_questions_is_not_saved(self):
        with mock.patch.object(ai, "_ask_claude", return_value=claude_reply(2)):
            with self.assertRaisesMessage(ai.GenerationError, "wrote 2 questions but the blueprint asks for 3"):
                ai.generate_final_paper(self.bp, self.controller)
        self.assertFalse(FinalPaper.objects.filter(questions__isnull=False).exists())

    def test_generation_runs_in_the_background_and_notifies(self):
        with mock.patch.object(ai.threading, "Thread") as thread:
            paper = ai.start_final_paper(self.bp, self.controller)
        thread.return_value.start.assert_called_once()
        self.assertTrue(paper.is_generating)
        with self.assertRaisesMessage(ai.GenerationError, "already writing"):
            ai.start_final_paper(self.bp, self.controller)

        with mock.patch.object(ai, "_ask_claude", return_value=claude_reply(3)):
            ai.write_final_paper(paper.pk, self.controller.pk)
        paper.refresh_from_db()
        self.assertFalse(paper.is_generating)
        self.assertEqual(paper.questions.count(), 3)
        self.assertTrue(Notification.objects.filter(recipient=self.controller, title__startswith="Final paper ready").exists())
        self.assertEqual(list(AuditEvent.objects.order_by("pk").values_list("action", flat=True)),
                         [AuditEvent.Action.FINAL_STARTED, AuditEvent.Action.FINAL_GENERATED])

    def test_failed_background_generation_is_shown(self):
        with mock.patch.object(ai.threading, "Thread"):
            paper = ai.start_final_paper(self.bp, self.controller)
        with mock.patch.object(ai, "_ask_claude", side_effect=ai.GenerationError("Claude is busy")):
            ai.write_final_paper(paper.pk, self.controller.pk)
        paper.refresh_from_db()
        self.assertEqual(paper.generation_error, "Claude is busy")
        self.assertFalse(paper.is_generating)


class DocumentUploadTests(BlueprintTestCase):
    def upload(self, name, content):
        self.client.force_login(self.controller)
        return self.client.post(reverse("admin:questionbank_knowledgesource_add"), {
            "subject": self.bp.subject.pk, "grade": self.bp.grade.pk, "material_type": "lesson_plan",
            "title": "Lesson 1", "chapter": "Chapter 1", "topic": "", "is_active": "on",
            "file": SimpleUploadedFile(name, content),
        })

    def test_word_file_text_is_extracted_and_sent_to_claude(self):
        doc = Document()
        doc.add_paragraph("Photosynthesis happens in chloroplasts.")
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Light"
        buffer = BytesIO()
        doc.save(buffer)
        self.upload("lesson.docx", buffer.getvalue())
        source = KnowledgeSource.objects.get(title="Lesson 1")
        self.assertIn("chloroplasts", source.content_text)
        self.assertIn("Light", source.content_text)
        block = ai._source_block(source)
        self.assertEqual(block["source"]["type"], "text")
        self.assertIn("chloroplasts", block["source"]["data"])
        self.assertEqual(block["title"], "Lesson plan: Lesson 1 · Chapter 1")

    def test_damaged_word_file_is_refused(self):
        response = self.upload("broken.docx", b"not really a word file")
        self.assertContains(response, "couldn")
        self.assertFalse(KnowledgeSource.objects.filter(title="Lesson 1").exists())


class QualityCheckTests(BlueprintTestCase):
    def generate_questions(self, texts):
        reply = {"questions": [{"text": t, "options": [], "answer": "A", "explanation": "", "source_reference": ""}
                               for t in texts]}
        with mock.patch.object(ai, "_ask_claude", return_value=reply):
            return ai.generate_questions(subject=self.bp.subject, grade=self.bp.grade, topic="",
                                         question_type="short", difficulty="easy", count=len(texts),
                                         language="en", user=self.controller)

    def test_duplicates_are_dropped_within_a_batch_and_flagged_against_the_bank(self):
        first = self.generate_questions(["What do plants need for photosynthesis?"])[0]
        created = self.generate_questions([
            "What do plants need for photosynthesis ?",   # same as the bank question
            "Name the green pigment in leaves.",
            "Name the green pigment in leaves",            # repeat within this batch
        ])
        self.assertEqual(len(created), 2)
        self.assertEqual(created[0].duplicate_of_id, first.pk)
        self.assertIsNone(created[1].duplicate_of_id)

    def test_final_paper_flags_copied_samples_repeats_and_ai_review_findings(self):
        item = self.bp.items.get(section="A")
        item.sample_question = "Which gas do plants release during photosynthesis?"
        item.save()
        reply = claude_reply(3)
        reply["questions"][0]["text"] = "Which gas do plants release during photosynthesis?"
        reply["questions"][2]["text"] = reply["questions"][1]["text"]
        review = {"reviews": [
            {"number": 1, "in_syllabus": True, "answer_correct": True, "issue": ""},
            {"number": 2, "in_syllabus": False, "answer_correct": False, "issue": "Chapter 9 isn't in the books."},
            {"number": 3, "in_syllabus": True, "answer_correct": True, "issue": ""},
        ]}
        with mock.patch.object(ai, "_ask_claude", side_effect=[reply, review]):
            paper = ai.generate_final_paper(self.bp, self.controller)
        notes = [q.check_notes for q in paper.questions.all()]
        self.assertIn("sample question", notes[0])
        self.assertIn("May not be covered", notes[1])
        self.assertIn("The answer may be wrong", notes[1])
        self.assertIn("Chapter 9", notes[1])
        self.assertIn("Very similar to question 2", notes[2])
        self.assertTrue(paper.check_summary.startswith("3 of 3 questions flagged"))

    def test_final_paper_is_saved_even_if_the_review_pass_fails(self):
        with mock.patch.object(ai, "_ask_claude", side_effect=[claude_reply(3), ai.GenerationError("busy")]):
            paper = ai.generate_final_paper(self.bp, self.controller)
        self.assertEqual(paper.questions.count(), 3)
        self.assertIn("couldn't run", paper.check_summary)


class MaterialChoiceTests(BlueprintTestCase):
    def setUp(self):
        super().setUp()
        self.book = KnowledgeSource.objects.get(title="Book")
        self.chapter = KnowledgeSource.objects.create(
            subject=self.bp.subject, grade=self.bp.grade, title="Chapter 5", chapter="Chapter 5",
            file=SimpleUploadedFile("ch5.txt", b"Respiration releases energy."))

    @staticmethod
    def sent_titles(ask):
        content = ask.call_args.args[1][0]["content"]
        return [block["title"] for block in content if block["type"] == "document"]

    def test_only_chosen_materials_reach_claude(self):
        with mock.patch.object(ai, "_ask_claude", return_value=claude_reply(1)) as ask:
            ai.generate_questions(subject=self.bp.subject, grade=self.bp.grade, topic="", question_type="short",
                                  difficulty="easy", count=1, language="en", user=self.controller,
                                  sources=[self.chapter])
        self.assertEqual(self.sent_titles(ask), ["Textbook: Chapter 5 · Chapter 5"])

    def test_blueprint_materials_limit_the_final_paper(self):
        self.bp.materials.add(self.chapter)
        with mock.patch.object(ai, "_ask_claude", return_value=claude_reply(3)) as ask:
            ai.generate_final_paper(self.bp, self.controller)
        self.assertEqual(self.sent_titles(ask), ["Textbook: Chapter 5 · Chapter 5"])

    def test_generate_page_lists_materials_with_their_subject_and_grade(self):
        self.client.force_login(self.controller)
        response = self.client.get(reverse("admin:questionbank_question_generate"))
        self.assertContains(response, f'data-pair="{self.bp.subject.pk}:{self.bp.grade.pk}"', count=2)

    def test_material_for_another_subject_is_refused(self):
        other = KnowledgeSource.objects.create(
            subject=Subject.objects.create(name="History"), grade=self.bp.grade, title="History book",
            file=SimpleUploadedFile("h.txt", b"Kings"))
        form = GenerateQuestionsForm(data={
            "subject": self.bp.subject.pk, "grade": self.bp.grade.pk, "question_type": "short",
            "difficulty": "easy", "count": 1, "language": "en", "sources": [other.pk]})
        self.assertFalse(form.is_valid())
        self.assertIn("isn't for Science / Class 8", form.non_field_errors()[0])


class ExportTests(BlueprintTestCase):
    def setUp(self):
        super().setUp()
        reply = claude_reply(3)
        reply["questions"][0].update(text="সালোকসংশ্লেষণ কী?", options=["ক", "খ", "গ", "ঘ"], answer="ক")
        with mock.patch.object(ai, "_ask_claude", return_value=reply):
            self.paper = ai.generate_final_paper(self.bp, self.controller)

    def download(self, fmt, user=None, answers=False):
        self.client.force_login(user or self.controller)
        url = reverse("admin:questionbank_final_export", args=[self.paper.pk, fmt])
        return self.client.get(url + ("?answers=1" if answers else ""))

    def test_unlocked_paper_cannot_be_downloaded(self):
        self.assertEqual(self.download("pdf").status_code, 302)

    def test_locked_paper_downloads_as_pdf_and_word(self):
        workflow.lock_final(self.paper, self.controller)
        pdf = self.download("pdf")
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertIn("science-mid-term-question-paper.pdf", pdf["Content-Disposition"])
        text = pymupdf.open(stream=pdf.content, filetype="pdf")[0].get_text()
        self.assertIn("Total marks: 7", text)
        self.assertIn("Question 1?", text)

        word = Document(BytesIO(self.download("docx", answers=True).content))
        body = "\n".join(p.text for p in word.paragraphs)
        self.assertIn("সালোকসংশ্লেষণ কী?", body)
        self.assertIn("Answer: ক", body)
        self.assertEqual(AuditEvent.objects.filter(action=AuditEvent.Action.FINAL_EXPORTED).count(), 2)

    def test_hod_cannot_download(self):
        workflow.lock_final(self.paper, self.controller)
        self.assertEqual(self.download("pdf", user=self.hod).status_code, 403)


@override_settings(MAKE_WEBHOOK_URL="", MAKE_MAILHOOK_ADDRESS="")
class ReminderTests(BlueprintTestCase):
    def setUp(self):
        super().setUp()
        User = get_user_model()
        self.teacher = User.objects.create_user("teacher", "teacher@example.com", "Pw-12345!", is_staff=True)
        self.teacher.user_permissions.add(Permission.objects.get(codename="add_question"))
        self.teacher = User.objects.get(pk=self.teacher.pk)
        self.bp.assigned_to, self.bp.status = self.teacher, QuestionBlueprint.Status.DRAFT
        self.bp.due_date = date.today()
        self.bp.save()

    def deadline(self, subject):
        return AcademicEvent.objects.create(title="Papers due", event_type=AcademicEvent.EventType.QUESTION_DEADLINE,
                                            start_date=date.today(), subject=subject)

    def test_deadline_reminders_only_for_their_subject(self):
        self.deadline(self.bp.subject)
        self.deadline(Subject.objects.create(name="History"))
        titles = [n.title for n in ensure_reminders(self.teacher)]
        self.assertEqual(sum("Question deadline" in t for t in titles), 1)

    def test_daily_command_emails_new_reminders_once(self):
        call_command("send_reminders", stdout=StringIO())
        self.assertEqual([m.to for m in mail.outbox], [["teacher@example.com"]])
        self.assertIn("Science mid-term", mail.outbox[0].body)
        call_command("send_reminders", stdout=StringIO())
        self.assertEqual(len(mail.outbox), 1)


class RoleTests(BlueprintTestCase):
    def setUp(self):
        super().setUp()
        call_command("setup_academics", stdout=StringIO())
        User = get_user_model()
        self.teacher = User.objects.create_user("teacher", password="Pw-12345!", is_staff=True)
        self.teacher.groups.add(Group.objects.get(name="Teacher / Subject Expert"))
        self.viewer = User.objects.create_user("viewer", password="Pw-12345!", is_staff=True)
        self.viewer.groups.add(Group.objects.get(name="Read-Only"))
        self.book = KnowledgeSource.objects.get()  # uploaded by nobody in particular

    def upload(self):
        return self.client.post(reverse("admin:questionbank_knowledgesource_add"), {
            "subject": self.bp.subject.pk, "grade": self.bp.grade.pk, "material_type": "syllabus",
            "title": "Science syllabus", "chapter": "", "topic": "", "is_active": "on",
            "file": SimpleUploadedFile("syllabus.txt", b"Unit 1: Cells"),
        })

    def test_teacher_uploads_and_edits_only_own_materials(self):
        self.client.force_login(self.teacher)
        self.upload()
        mine = KnowledgeSource.objects.get(title="Science syllabus")
        self.assertEqual(mine.uploaded_by, self.teacher)
        self.assertEqual(mine.label, "Syllabus: Science syllabus")
        change = "admin:questionbank_knowledgesource_change"
        self.assertContains(self.client.get(reverse(change, args=[mine.pk])), 'name="_save"')
        self.assertNotContains(self.client.get(reverse(change, args=[self.book.pk])), 'name="_save"')

    def test_read_only_user_can_view_but_not_add(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("admin:questionbank_knowledgesource_changelist")).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:questionbank_question_changelist")).status_code, 200)
        self.assertEqual(self.upload().status_code, 403)
        self.assertEqual(self.client.get(reverse("admin:questionbank_questionblueprint_changelist")).status_code, 403)


class ReviewTests(BlueprintTestCase):
    def test_hod_cannot_approve_own_blueprint(self):
        self.bp.status = QuestionBlueprint.Status.IN_REVIEW
        self.bp.save()
        with self.assertRaisesMessage(workflow.WorkflowError, "your own work"):
            workflow.approve(self.bp, self.hod)
        workflow.approve(self.bp, self.controller)  # someone else can
        self.assertEqual(self.bp.status, QuestionBlueprint.Status.APPROVED)

    def test_super_admin_cannot_approve_own_blueprint_either(self):
        self.bp.status, self.bp.assigned_to = QuestionBlueprint.Status.IN_REVIEW, self.controller
        self.bp.save()
        with self.assertRaisesMessage(workflow.WorkflowError, "your own work"):
            workflow.approve(self.bp, self.controller)


class QuestionApprovalTests(BlueprintTestCase):
    def setUp(self):
        super().setUp()
        call_command("setup_academics", stdout=StringIO())
        User = get_user_model()
        self.teacher = User.objects.create_user("teacher", password="Pw-12345!", is_staff=True)
        self.teacher.groups.add(Group.objects.get(name="Teacher / Subject Expert"))
        self.head = User.objects.create_user("head", password="Pw-12345!", is_staff=True)
        self.head.groups.add(Group.objects.get(name="Head of Department"))
        make = lambda by, text: Question.objects.create(subject=self.bp.subject, grade=self.bp.grade,
                                                        question_type="short", text=text, answer="A", created_by=by)
        self.teachers_q, self.heads_q = make(self.teacher, "Teacher's question"), make(self.head, "HOD's question")

    def act(self, user, action, *questions):
        self.client.force_login(user)
        return self.client.post(reverse("admin:questionbank_question_changelist"),
                                {"action": action, "_selected_action": [q.pk for q in questions]})

    def status(self, q):
        q.refresh_from_db()
        return q.status

    def test_teacher_cannot_approve_questions(self):
        self.act(self.teacher, "approve", self.teachers_q)
        self.assertEqual(self.status(self.teachers_q), Question.Status.DRAFT)

    def test_teacher_cannot_approve_by_editing_status(self):
        self.client.force_login(self.teacher)
        url = reverse("admin:questionbank_question_change", args=[self.teachers_q.pk])
        self.assertNotContains(self.client.get(url), 'name="status"')

    def test_hod_approves_others_but_not_own(self):
        self.act(self.head, "approve", self.teachers_q, self.heads_q)
        self.assertEqual(self.status(self.teachers_q), Question.Status.APPROVED)
        self.assertEqual(self.status(self.heads_q), Question.Status.DRAFT)

    def test_editing_an_approved_question_sends_it_back_to_draft(self):
        self.act(self.head, "approve", self.teachers_q)
        self.client.force_login(self.teacher)
        self.client.post(reverse("admin:questionbank_question_change", args=[self.teachers_q.pk]), {
            "subject": self.bp.subject.pk, "grade": self.bp.grade.pk, "topic": "", "question_type": "short",
            "difficulty": "medium", "text": "Teacher's question, reworded", "options": "[]", "answer": "A",
            "explanation": "", "source_reference": ""})
        self.assertEqual(self.status(self.teachers_q), Question.Status.DRAFT)
