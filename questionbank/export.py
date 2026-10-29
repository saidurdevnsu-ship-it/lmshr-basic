"""Export a locked final paper (or its answer key) as PDF or Word. Exam controller only; see FinalPaperAdmin."""

import io

import pymupdf
from django.conf import settings
from django.template.loader import render_to_string
from django.utils.text import slugify
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

# Windows / Word font for Bangla; Word falls back sensibly where it isn't installed.
BANGLA_FONT = "Nirmala UI"


def paper_context(paper, answers):
    """Everything the print page and the exports show: the blueprint details and questions grouped by section."""
    questions = list(paper.questions.all())
    sections = []
    for q in questions:
        if not sections or sections[-1]["name"] != q.section:
            sections.append({"name": q.section, "questions": []})
        sections[-1]["questions"].append(q)
    return {
        "paper": paper,
        "bp": paper.blueprint,
        "sections": sections,
        "total_marks": sum(q.marks for q in questions),
        "answers": answers,
        "company_name": settings.HR_COMPANY_NAME,
    }


def filename(paper, answers, extension):
    name = slugify(paper.blueprint.title) or f"final-paper-{paper.pk}"
    return f"{name}-{'answer-key' if answers else 'question-paper'}.{extension}"


def export_pdf(paper, answers=False):
    html = render_to_string("questionbank/final_export.html", paper_context(paper, answers))
    story = pymupdf.Story(html=html)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page = pymupdf.paper_rect("a4")
    body = page + (56, 56, -56, -64)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(body)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Page numbers, added once the page count is known.
    doc = pymupdf.open(stream=buffer.getvalue(), filetype="pdf")
    for number, pdf_page in enumerate(doc, 1):
        pdf_page.insert_text((page.width / 2 - 30, page.height - 32), f"Page {number} of {doc.page_count}",
                             fontsize=8, color=(0.4, 0.4, 0.4))
    return doc.tobytes(garbage=3, deflate=True)


def export_docx(paper, answers=False):
    ctx = paper_context(paper, answers)
    bp = ctx["bp"]
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), BANGLA_FONT)

    def centered(text, size, bold=False):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        return p

    centered(ctx["company_name"], 16, bold=True)
    centered(bp.exam.title if bp.exam else "Examination", 11)
    centered("Answer key · confidential" if answers else bp.title, 14, bold=True)
    centered(f"{bp.subject} · {bp.grade}", 11)
    centered(f"Time allowed: {bp.duration_minutes} minutes        Total marks: {ctx['total_marks']}", 11)
    if paper.status != paper.Status.LOCKED:
        centered("NOT LOCKED. This paper can still change.", 10, bold=True)

    if not answers and bp.instructions.strip():
        p = document.add_paragraph()
        p.add_run("Instructions: ").bold = True
        p.add_run(bp.instructions.strip())

    for section in ctx["sections"]:
        if section["name"]:
            document.add_heading(section["name"], level=2)
        for q in section["questions"]:
            p = document.add_paragraph()
            p.add_run(f"{q.order}. ").bold = True
            p.add_run(q.text)
            p.add_run(f"  [{q.marks}]").bold = True
            for letter, option in zip("ABCDEFGH", q.options):
                document.add_paragraph(f"{letter}. {option}").paragraph_format.left_indent = Pt(24)
            if answers:
                a = document.add_paragraph()
                a.paragraph_format.left_indent = Pt(24)
                a.add_run("Answer: ").bold = True
                a.add_run(q.answer)
                if q.explanation:
                    a.add_run(f"\n{q.explanation}").italic = True
                if q.source_reference:
                    a.add_run(f"\nSource: {q.source_reference}").font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
