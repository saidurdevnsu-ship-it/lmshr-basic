"""Read the text out of uploaded materials: PDF (PyMuPDF), Word (python-docx), plain text and Markdown."""

from io import BytesIO

import pymupdf
from docx import Document

ALLOWED_EXTENSIONS = ["pdf", "docx", "txt", "md"]


class UnreadableDocument(Exception):
    pass


def extension(name):
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def extract_text(file, name=None):
    """All the text in an uploaded or stored file. Scanned PDFs have no text layer and give an empty string."""
    name = name or file.name
    file.seek(0)
    data = file.read()
    file.seek(0)
    try:
        ext = extension(name)
        if ext == "pdf":
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                return "\n".join(page.get_text() for page in doc).strip()
        if ext == "docx":
            return _docx_text(Document(BytesIO(data)))
        return data.decode("utf-8", errors="replace").strip()
    except Exception as exc:
        raise UnreadableDocument(f"This file couldn't be read ({exc.__class__.__name__}). Is it damaged?") from exc


def _docx_text(document):
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())
